from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.ext.asyncio import AsyncSession

from app.enums import ReconciliationState, UserRole, UserStatus
from app.errors import DomainError
from app.models import ReferenceOrder, ReferenceUpload, User
from app.repositories.reconciliation import ReconciliationRepository
from app.validators import clean_display_name, normalize_name, validate_order_number

MAX_REFERENCE_FILE_BYTES = 10 * 1024 * 1024
SCIENTIFIC_NUMBER_RE = re.compile(r"^\d+(?:[.,]\d+)?[Ee][+-]?\d+$")


@dataclass(slots=True, frozen=True)
class ParsedReferenceOrder:
    raw: str
    normalized: str
    supplier_name: str | None


@dataclass(slots=True, frozen=True)
class ParsedReferenceWorkbook:
    rows: tuple[ParsedReferenceOrder, ...]
    warning_count: int


@dataclass(slots=True, frozen=True)
class ReconciliationMatch:
    state: ReconciliationState
    supplier_name: str | None
    upload_id: int | None


def _order_cell_to_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        raise DomainError("Пустой или некорректный номер заказа.")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not value.is_integer() or abs(value) >= 1_000_000_000_000_000:
            raise DomainError("Числовой номер заказа невозможно прочитать без потери точности.")
        return f"{value:.0f}"
    text = str(value).strip()
    if SCIENTIFIC_NUMBER_RE.fullmatch(text):
        try:
            number = Decimal(text.replace(",", "."))
        except InvalidOperation as exc:
            raise DomainError("Не удалось прочитать номер заказа.") from exc
        if number != number.to_integral_value():
            raise DomainError("Номер заказа в научной записи должен быть целым.")
        text = format(number, "f")
    return text


def parse_reference_workbook(content: bytes) -> ParsedReferenceWorkbook:
    if not content:
        raise DomainError("Файл пустой.")
    if len(content) > MAX_REFERENCE_FILE_BYTES:
        raise DomainError("Файл слишком большой. Максимальный размер — 10 МБ.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise DomainError("Не удалось открыть Excel-файл. Нужен корректный .xlsx.") from exc
    try:
        sheet = workbook.active
        headers = [
            normalize_name(str(sheet.cell(1, column).value or ""))
            for column in (1, 2)
        ]
        if headers[0] not in {"заказ", "номер заказа"} or headers[1] != "поставщик":
            raise DomainError(
                "В первой строке должны быть колонки «Заказ» и «Поставщик»."
            )
        for column in range(3, sheet.max_column + 1):
            if any(sheet.cell(row, column).value not in (None, "") for row in range(1, sheet.max_row + 1)):
                raise DomainError("В файле должны быть только две заполненные колонки.")

        parsed: dict[str, ParsedReferenceOrder] = {}
        supplier_keys: dict[str, str] = {}
        warning_count = 0
        errors: list[str] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=2, values_only=True), start=2
        ):
            order_value, supplier_value = values
            if order_value in (None, "") and supplier_value in (None, ""):
                continue
            try:
                raw_value = _order_cell_to_text(order_value)
                raw, normalized = validate_order_number(raw_value)
            except DomainError as exc:
                errors.append(f"Строка {row_number}: {exc}")
                continue
            supplier_name: str | None = None
            if supplier_value not in (None, ""):
                try:
                    supplier_name = clean_display_name(str(supplier_value))
                except DomainError as exc:
                    errors.append(f"Строка {row_number}: {exc}")
                    continue
            else:
                warning_count += 1
            supplier_key = normalize_name(supplier_name or "")
            existing = parsed.get(normalized)
            if existing is not None:
                if supplier_keys[normalized] != supplier_key:
                    errors.append(
                        f"Строка {row_number}: у заказа {raw} указаны разные поставщики."
                    )
                continue
            parsed[normalized] = ParsedReferenceOrder(raw, normalized, supplier_name)
            supplier_keys[normalized] = supplier_key
        if errors:
            preview = "\n".join(errors[:10])
            suffix = "\n…" if len(errors) > 10 else ""
            raise DomainError(f"Файл содержит ошибки:\n{preview}{suffix}")
        if not parsed:
            raise DomainError("В файле нет заказов для загрузки.")
        return ParsedReferenceWorkbook(tuple(parsed.values()), warning_count)
    finally:
        workbook.close()


class ReconciliationService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.repository = ReconciliationRepository(session)

    @staticmethod
    def _validate_admin(user: User | None) -> User:
        if user is None or user.status is not UserStatus.ACTIVE or user.role is not UserRole.ADMIN:
            raise DomainError("Загружать файл сверки может только администратор.")
        return user

    async def import_xlsx(
        self,
        *,
        content: bytes,
        filename: str,
        uploaded_by: User | None,
    ) -> ReferenceUpload:
        admin = self._validate_admin(uploaded_by)
        if not filename.lower().endswith(".xlsx"):
            raise DomainError("Нужен файл в формате .xlsx.")
        parsed = parse_reference_workbook(content)
        await self.repository.get_active_upload(for_update=True)
        await self.repository.deactivate_all()
        upload = ReferenceUpload(
            uploaded_by_id=admin.id,
            original_filename=filename[:255],
            sha256=hashlib.sha256(content).hexdigest(),
            row_count=len(parsed.rows),
            warning_count=parsed.warning_count,
            active=True,
        )
        upload.rows = [
            ReferenceOrder(
                order_number_raw=row.raw,
                order_number_normalized=row.normalized,
                supplier_name=row.supplier_name,
            )
            for row in parsed.rows
        ]
        self.repository.add(upload)
        await self.repository.flush()
        return upload

    async def get_active_upload(self) -> ReferenceUpload | None:
        return await self.repository.get_active_upload()

    async def resolve(self, normalized_order: str) -> ReconciliationMatch:
        upload = await self.repository.get_active_upload()
        if upload is None:
            return ReconciliationMatch(ReconciliationState.NO_FILE, None, None)
        row = await self.repository.find_order(upload.id, normalized_order)
        if row is None:
            return ReconciliationMatch(ReconciliationState.MISSING, None, upload.id)
        return ReconciliationMatch(
            ReconciliationState.FOUND,
            row.supplier_name,
            upload.id,
        )
