from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.enums import InspectionWorkType
from app.services.exports import ExportRow

HEADERS = (
    "№",
    "Фамилия",
    "Дата инспекции",
    "Номер заказа",
    "Поставщик",
    "статус",
    "Название проекта",
    "Ставка за инспекцию",
    "Комментарии",
    "дата добавления строки",
    "дата внесения изменений",
    "статус обработки",
)

COLUMN_WIDTHS = (
    6.44140625,
    22.0,
    14.77734375,
    17.77734375,
    21.5546875,
    15.77734375,
    16.33203125,
    16.33203125,
    13.0,
    22.6640625,
    23.109375,
    13.0,
)
HEADER_BLUE = "44B3E1"
HEADER_YELLOW = "FFC000"
ROW_EVEN = "F3F9FC"
PROCESSED_GREEN = "E2F0D9"
UNPROCESSED_YELLOW = "FFF2CC"
DUPLICATE_RED = "C00000"
THIN_GRAY = Side(style="thin", color="A6A6A6")


def _localize(value: datetime, timezone: ZoneInfo) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone).replace(tzinfo=None)
    return value.astimezone(timezone).replace(tzinfo=None)


def _populate_sheet(
    sheet: Worksheet, rows: list[ExportRow], *, timezone: ZoneInfo
) -> None:
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:L1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.sheet_view.zoomScale = 85

    for column_index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for cell in sheet[1]:
        cell.fill = PatternFill(
            "solid", fgColor=HEADER_YELLOW if cell.column == 10 else HEADER_BLUE
        )
        cell.font = Font(name="Aptos Narrow", size=11, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY, bottom=THIN_GRAY)
    sheet.row_dimensions[1].height = 34

    group_numbers: dict[int, int] = {}
    for item in rows:
        if item.inspection_id not in group_numbers:
            group_numbers[item.inspection_id] = len(group_numbers) + 1
        group_number = group_numbers[item.inspection_id]
        processed_text = (
            f"Выгружено {_localize(item.processed_at, timezone):%d.%m.%Y %H:%M}"
            if item.processed_at
            else "Не выгружено"
        )
        values = (
            group_number,
            item.expert_name,
            item.inspection_date.strftime("%d.%m.%Y"),
            item.order_number,
            item.supplier_name or "",
            item.status_name,
            item.project_name,
            float(item.rate),
            item.comment or "",
            (
                _localize(item.created_at, timezone).strftime("%d.%m.%Y %H:%M")
                if item.created_at
                else ""
            ),
            (
                _localize(item.modified_at, timezone).strftime("%d.%m.%Y %H:%M")
                if item.modified_at
                else None
            ),
            processed_text,
        )
        sheet.append(values)
        excel_row = sheet.max_row
        sheet.row_dimensions[excel_row].height = 24
        for cell in sheet[excel_row]:
            cell.font = Font(name="Aptos Narrow", size=10, color="000000")
            cell.alignment = Alignment(
                vertical="center", wrap_text=cell.column in {5, 6, 9, 12}
            )
            cell.border = Border(
                top=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY, bottom=THIN_GRAY
            )
            if group_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ROW_EVEN)

        sheet.cell(excel_row, 1).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(excel_row, 3).number_format = "@"
        sheet.cell(excel_row, 3).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(excel_row, 4).number_format = "@"
        sheet.cell(excel_row, 8).number_format = "#,##0.00"
        sheet.cell(excel_row, 10).number_format = "@"
        sheet.cell(excel_row, 11).number_format = "@"
        sheet.cell(excel_row, 12).fill = PatternFill(
            "solid", fgColor=PROCESSED_GREEN if item.processed_at else UNPROCESSED_YELLOW
        )
        if item.is_duplicate:
            sheet.cell(excel_row, 4).font = Font(
                name="Aptos Narrow", size=10, color=DUPLICATE_RED, bold=True
            )

    sheet.auto_filter.ref = f"A1:L{max(sheet.max_row, 1)}"


def build_inspection_workbook(rows: list[ExportRow], *, timezone: ZoneInfo) -> bytes:
    workbook = Workbook(iso_dates=True)
    workbook.properties.title = "Отчёт по инспекциям"
    workbook.properties.subject = "Выгрузка Telegram-бота учёта инспекций"
    workbook.properties.creator = "Inspection Telegram Bot"

    role_sheets = (
        ("Специалисты", InspectionWorkType.SPECIALIST),
        ("Эксперты", InspectionWorkType.EXPERT),
    )
    for index, (title, work_type) in enumerate(role_sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        _populate_sheet(
            sheet,
            [item for item in rows if item.work_type is work_type],
            timezone=timezone,
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
