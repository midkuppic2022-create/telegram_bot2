from dataclasses import replace
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.enums import ExportGroup, InspectionWorkType
from app.exporting.xlsx import COLUMN_WIDTHS, HEADERS, build_inspection_workbook
from app.services.exports import ExportRow


def make_row(
    inspection_id: int,
    group: ExportGroup,
    order: str,
    *,
    work_type: InspectionWorkType = InspectionWorkType.EXPERT,
    duplicate: bool = False,
    processed: bool = False,
) -> ExportRow:
    return ExportRow(
        inspection_id=inspection_id,
        group_number=inspection_id,
        position=0,
        expert_name="Петросян Артур",
        inspection_date=date(2026, 7, 14),
        order_number=order,
        normalized_order=order,
        normalized_orders=(order,),
        work_type=work_type,
        supplier_name="Гранада ООО",
        status_name="осмотр / стоп-отгрузка",
        project_name=group.value,
        export_group=group,
        rate=Decimal("1500.00"),
        comment="Без замечаний",
        created_at=datetime(2026, 7, 14, 10, 30, tzinfo=UTC),
        modified_at=None,
        processed_at=datetime(2026, 7, 14, 12, 0, tzinfo=UTC) if processed else None,
        is_duplicate=duplicate,
    )


def test_workbook_structure_and_styles() -> None:
    content = build_inspection_workbook(
        [
            make_row(1, ExportGroup.AK, "YUG1", duplicate=True, processed=True),
            make_row(
                2,
                ExportGroup.SV,
                "SV1",
                work_type=InspectionWorkType.SPECIALIST,
            ),
            make_row(3, ExportGroup.X5, "X51"),
        ],
        timezone=ZoneInfo("Europe/Moscow"),
    )
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Специалисты", "Эксперты"]

    sheet = workbook["Эксперты"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:L3"
    assert sheet.max_row == 3
    assert [sheet.cell(row, 1).value for row in range(2, 4)] == [1, 2]
    assert sheet["C2"].value == "14.07.2026"
    assert sheet["C2"].number_format == "@"
    assert sheet["D2"].number_format == "@"
    assert sheet["H2"].number_format == "#,##0.00"
    assert sheet["K2"].value is None
    assert sheet["J2"].value == "14.07.2026 13:30"
    assert sheet["J2"].number_format == "@"
    assert (
        tuple(sheet.column_dimensions[column].width for column in "ABCDEFGHIJKL")
        == COLUMN_WIDTHS
    )

    assert sheet["D2"].font.color.rgb == "00C00000"
    assert sheet["L2"].value.startswith("Выгружено")
    assert sheet["L3"].value == "Не выгружено"
    assert sheet["J1"].fill.fgColor.rgb == "00FFC000"

    specialist_sheet = workbook["Специалисты"]
    assert specialist_sheet.max_row == 2
    assert specialist_sheet["D2"].value == "SV1"


def test_empty_workbook_has_two_role_sheets_with_headers() -> None:
    content = build_inspection_workbook([], timezone=ZoneInfo("Europe/Moscow"))
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Специалисты", "Эксперты"]
    assert workbook["Специалисты"].max_row == 1
    assert workbook["Эксперты"].max_row == 1


def test_naive_processed_timestamp_is_supported() -> None:
    row = replace(
        make_row(1, ExportGroup.AK, "NAIVE1", processed=True),
        processed_at=datetime(2026, 7, 14, 12, 0),
    )
    content = build_inspection_workbook([row], timezone=ZoneInfo("Europe/Moscow"))
    workbook = load_workbook(BytesIO(content))
    assert workbook["Эксперты"]["L2"].value == "Выгружено 14.07.2026 12:00"


def test_modified_timestamp_is_written() -> None:
    row = replace(
        make_row(1, ExportGroup.AK, "EDITED1"),
        modified_at=datetime(2026, 7, 15, 8, 45, tzinfo=UTC),
    )
    content = build_inspection_workbook([row], timezone=ZoneInfo("Europe/Moscow"))
    workbook = load_workbook(BytesIO(content))

    assert workbook["Эксперты"]["K2"].value == "15.07.2026 11:45"
    assert workbook["Эксперты"]["K2"].number_format == "@"


def test_fuel_column_is_removed_from_excel() -> None:
    row = make_row(1, ExportGroup.AK, "SPECIALIST1")
    content = build_inspection_workbook([row], timezone=ZoneInfo("Europe/Moscow"))
    workbook = load_workbook(BytesIO(content))

    assert "ГСМ" not in HEADERS
    assert workbook["Эксперты"]["H2"].value == 1500


def test_additional_order_is_a_separate_zero_rate_row() -> None:
    first = replace(
        make_row(7, ExportGroup.X5, "MAIN7", processed=True),
        group_number=3,
        normalized_orders=("MAIN7", "EXTRA7"),
    )
    second = replace(
        first,
        position=1,
        order_number="EXTRA7",
        normalized_order="EXTRA7",
        supplier_name="Другой поставщик",
        status_name="-",
        rate=Decimal("0.00"),
        comment=None,
        created_at=first.created_at,
        modified_at=first.modified_at,
        processed_at=first.processed_at,
    )

    content = build_inspection_workbook([first, second], timezone=ZoneInfo("Europe/Moscow"))
    sheet = load_workbook(BytesIO(content))["Эксперты"]

    assert [sheet["A2"].value, sheet["A3"].value] == [1, 1]
    assert [sheet["D2"].value, sheet["D3"].value] == ["MAIN7", "EXTRA7"]
    assert [sheet["E2"].value, sheet["E3"].value] == [
        "Гранада ООО",
        "Другой поставщик",
    ]
    assert sheet["F3"].value == "-"
    assert sheet["H3"].value == 0
    assert sheet["H3"].number_format == "#,##0.00"
    assert sheet["I3"].value is None
    assert sheet["J3"].value == "14.07.2026 13:30"
    assert sheet["L3"].value == "Выгружено 14.07.2026 15:00"
